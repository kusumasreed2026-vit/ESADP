import os
import glob
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


INPUT_FOLDER = r"D:\Shrihi\M.Tech  VITCC\Data Engineering\Use Case\CSV files"

OUTPUT_FOLDER = r"D:\Shrihi\M.Tech  VITCC\Data Engineering\Use Case\data_profiling_reports"

WORD_OUTPUT = os.path.join(
    OUTPUT_FOLDER,
    "Sprint2_Data_Quality_Report.docx"
)

EXCEL_OUTPUT = os.path.join(
    OUTPUT_FOLDER,
    "Sprint2_Data_Quality_Findings.xlsx"
)


os.makedirs(OUTPUT_FOLDER, exist_ok=True)


def load_csv(filepath):

    try:
        return pd.read_csv(
            filepath,
            low_memory=False
        )

    except UnicodeDecodeError:

        return pd.read_csv(
            filepath,
            encoding="latin1",
            low_memory=False
        )


def add_finding(
    findings,
    dimension,
    dataset,
    column,
    issue,
    count,
    percentage,
    severity,
    treatment
):

    findings.append({
        "Data Quality Dimension": dimension,
        "Dataset": dataset,
        "Column": column,
        "Issue": issue,
        "Count": int(count),
        "Percentage": round(float(percentage), 2),
        "Severity": severity,
        "Recommended Treatment": treatment
    })


def add_table(document, dataframe, max_rows=40):

    if dataframe.empty:

        document.add_paragraph(
            "No findings identified for this section."
        )

        return

    display_df = dataframe.head(max_rows)

    table = document.add_table(
        rows=1,
        cols=len(display_df.columns)
    )

    table.style = "Table Grid"

    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header = table.rows[0].cells

    for i, column in enumerate(display_df.columns):

        header[i].text = str(column)

        for paragraph in header[i].paragraphs:

            for run in paragraph.runs:

                run.bold = True

                run.font.size = Pt(8)

        tcPr = header[i]._tc.get_or_add_tcPr()

        shd = OxmlElement("w:shd")

        shd.set(
            qn("w:fill"),
            "D9EAF7"
        )

        tcPr.append(shd)

    for _, row in display_df.iterrows():

        cells = table.add_row().cells

        for i, value in enumerate(row):

            if pd.isna(value):

                value = ""

            cells[i].text = str(value)

            for paragraph in cells[i].paragraphs:

                for run in paragraph.runs:

                    run.font.size = Pt(8)

    if len(dataframe) > max_rows:

        document.add_paragraph(
            f"Showing first {max_rows} findings. "
            f"Complete findings are available in the Excel report."
        )


print("=" * 75)
print("SPRINT 2 - OLIST DATA QUALITY ANALYSIS")
print("=" * 75)


print("\nDiscovering source CSV files...")


csv_files = sorted(
    glob.glob(
        os.path.join(
            INPUT_FOLDER,
            "**",
            "*.csv"
        ),
        recursive=True
    )
)


print(
    f"CSV files discovered: {len(csv_files)}"
)


if not csv_files:

    raise FileNotFoundError(
        "No CSV files were found. Check INPUT_FOLDER."
    )


datasets = {}


for filepath in csv_files:

    filename = os.path.basename(filepath)

    print(
        f"Loading: {filename}"
    )

    try:

        datasets[filename] = load_csv(
            filepath
        )

    except Exception as e:

        print(
            f"ERROR loading {filename}: {e}"
        )


print(
    f"\nDatasets loaded successfully: {len(datasets)}"
)


findings = []

summary_rows = []


print("\n" + "-" * 75)
print("1. COMPLETENESS AND GENERAL PROFILING CHECKS")
print("-" * 75)


for filename, df in datasets.items():

    row_count = len(df)

    column_count = len(df.columns)

    missing_cells = int(
        df.isna().sum().sum()
    )

    missing_columns = int(
        (df.isna().sum() > 0).sum()
    )

    duplicate_rows = int(
        df.duplicated().sum()
    )

    summary_rows.append({

        "Dataset": filename,

        "Rows": row_count,

        "Columns": column_count,

        "Missing Cells": missing_cells,

        "Columns With Missing Values": missing_columns,

        "Duplicate Full Rows": duplicate_rows

    })

    for column in df.columns:

        missing_count = int(
            df[column].isna().sum()
        )

        if missing_count == 0:

            continue

        missing_percentage = (
            missing_count /
            row_count *
            100
            if row_count > 0
            else 0
        )

        if missing_percentage >= 50:

            severity = "High"

        elif missing_percentage >= 5:

            severity = "Medium"

        else:

            severity = "Low"

        add_finding(

            findings,

            "Completeness",

            filename,

            column,

            "Missing values detected",

            missing_count,

            missing_percentage,

            severity,

            "Investigate the business meaning and retain NULL, impute, or assign a suitable default value according to the field."

        )


print("Completeness checks completed.")


print("\n" + "-" * 75)
print("2. KEY UNIQUENESS CHECKS")
print("-" * 75)


single_key_checks = [

    (
        "olist_orders_dataset.csv",
        "order_id"
    ),

    (
        "olist_products_dataset.csv",
        "product_id"
    ),

    (
        "olist_customers_dataset.csv",
        "customer_id"
    ),

    (
        "olist_sellers_dataset.csv",
        "seller_id"
    )

]


for filename, column in single_key_checks:

    if filename not in datasets:

        continue

    df = datasets[filename]

    if column not in df.columns:

        continue

    row_count = len(df)

    missing_count = int(
        df[column].isna().sum()
    )

    if missing_count > 0:

        percentage = (
            missing_count /
            row_count *
            100
            if row_count > 0
            else 0
        )

        add_finding(

            findings,

            "Uniqueness",

            filename,

            column,

            "Missing key identifier",

            missing_count,

            percentage,

            "High",

            "Investigate and resolve missing key identifiers before warehouse loading."

        )

    duplicate_count = int(
        df[column].duplicated(
            keep=False
        ).sum()
    )

    if duplicate_count > 0:

        percentage = (
            duplicate_count /
            row_count *
            100
            if row_count > 0
            else 0
        )

        add_finding(

            findings,

            "Uniqueness",

            filename,

            column,

            "Duplicate key identifiers detected",

            duplicate_count,

            percentage,

            "High",

            "Investigate duplicate key values and determine whether they represent true duplicates or source-system issues."

        )


print(
    "Single-column key checks completed."
)


print("\nChecking order-item composite key...")


order_items_filename = (
    "olist_order_items_dataset.csv"
)


if order_items_filename in datasets:

    df = datasets[
        order_items_filename
    ]

    required_columns = [
        "order_id",
        "order_item_id"
    ]

    if all(
        column in df.columns
        for column in required_columns
    ):

        composite_duplicates = int(

            df.duplicated(
                subset=[
                    "order_id",
                    "order_item_id"
                ],
                keep=False
            ).sum()

        )

        percentage = (

            composite_duplicates /
            len(df) *
            100
            if len(df) > 0
            else 0

        )

        if composite_duplicates > 0:

            add_finding(

                findings,

                "Uniqueness",

                order_items_filename,

                "order_id + order_item_id",

                "Duplicate order-item composite keys detected",

                composite_duplicates,

                percentage,

                "High",

                "Investigate repeated order-item combinations before warehouse loading."

            )

        else:

            print(
                "No duplicate order-item composite keys found."
            )


print(
    "Order-item composite key check completed."
)


print("\n" + "-" * 75)
print("3. ORDER ITEM VALIDITY CHECKS")
print("-" * 75)


if order_items_filename in datasets:

    df = datasets[
        order_items_filename
    ]

    if "price" in df.columns:

        negative_price = int(
            (
                pd.to_numeric(
                    df["price"],
                    errors="coerce"
                )
                < 0
            ).sum()
        )

        if negative_price > 0:

            percentage = (
                negative_price /
                len(df) *
                100
            )

            add_finding(

                findings,

                "Validity",

                order_items_filename,

                "price",

                "Negative price values detected",

                negative_price,

                percentage,

                "High",

                "Validate negative prices and correct or exclude invalid transaction records according to business rules."

            )


    if "freight_value" in df.columns:

        negative_freight = int(
            (
                pd.to_numeric(
                    df["freight_value"],
                    errors="coerce"
                )
                < 0
            ).sum()
        )

        if negative_freight > 0:

            percentage = (
                negative_freight /
                len(df) *
                100
            )

            add_finding(

                findings,

                "Validity",

                order_items_filename,

                "freight_value",

                "Negative freight values detected",

                negative_freight,

                percentage,

                "Medium",

                "Validate negative freight values against source-system business rules."

            )


    if "quantity" in df.columns:

        quantity_numeric = pd.to_numeric(
            df["quantity"],
            errors="coerce"
        )

        invalid_quantity = int(
            (
                quantity_numeric <= 0
            ).sum()
        )

        if invalid_quantity > 0:

            percentage = (
                invalid_quantity /
                len(df) *
                100
            )

            add_finding(

                findings,

                "Validity",

                order_items_filename,

                "quantity",

                "Non-positive quantity detected",

                invalid_quantity,

                percentage,

                "High",

                "Validate quantity values and correct invalid transaction records."

            )

    else:

        print(
            "Quantity column not present. "
            "Quantity validation skipped."
        )


print(
    "Order-item validity checks completed."
)


print("\n" + "-" * 75)
print("4. PRODUCT MASTER QUALITY CHECKS")
print("-" * 75)


products_filename = (
    "olist_products_dataset.csv"
)


if products_filename in datasets:

    df = datasets[
        products_filename
    ]

    product_measurements = [

        "product_weight_g",

        "product_length_cm",

        "product_height_cm",

        "product_width_cm"

    ]

    for column in product_measurements:

        if column not in df.columns:

            continue

        numeric_values = pd.to_numeric(
            df[column],
            errors="coerce"
        )

        invalid_values = int(
            (
                numeric_values < 0
            ).sum()
        )

        if invalid_values > 0:

            percentage = (
                invalid_values /
                len(df) *
                100
            )

            add_finding(

                findings,

                "Validity",

                products_filename,

                column,

                "Negative product measurement detected",

                invalid_values,

                percentage,

                "Medium",

                "Validate product measurements and correct invalid source values."

            )


    important_product_columns = [

        "product_category_name",

        "product_name_lenght",

        "product_description_lenght",

        "product_photos_qty",

        "product_weight_g",

        "product_length_cm",

        "product_height_cm",

        "product_width_cm"

    ]


    for column in important_product_columns:

        if column not in df.columns:

            continue

        missing = int(
            df[column].isna().sum()
        )

        if missing == 0:

            continue

        percentage = (
            missing /
            len(df) *
            100
        )

        severity = (

            "Medium"
            if percentage >= 5
            else "Low"

        )

        add_finding(

            findings,

            "Product Completeness",

            products_filename,

            column,

            "Missing product attribute",

            missing,

            percentage,

            severity,

            "Apply category mapping, NULL retention, or documented imputation according to the product-master business rule."

        )


print(
    "Product quality checks completed."
)


print("\n" + "-" * 75)
print("5. REFERENTIAL INTEGRITY CHECKS")
print("-" * 75)


orders_filename = (
    "olist_orders_dataset.csv"
)

payments_filename = (
    "olist_order_payments_dataset.csv"
)

sellers_filename = (
    "olist_sellers_dataset.csv"
)


orders = datasets.get(
    orders_filename
)

order_items = datasets.get(
    order_items_filename
)

payments = datasets.get(
    payments_filename
)

products = datasets.get(
    products_filename
)

sellers = datasets.get(
    sellers_filename
)


if orders is not None:

    order_ids = set(
        orders[
            "order_id"
        ].dropna()
    )

else:

    order_ids = set()


if order_items is not None:

    item_order_ids = set(
        order_items[
            "order_id"
        ].dropna()
    )

    orphan_items = int(

        (
            ~order_items[
                "order_id"
            ].isin(order_ids)
        ).sum()

    )

    if orphan_items > 0:

        percentage = (

            orphan_items /
            len(order_items) *
            100

        )

        add_finding(

            findings,

            "Referential Integrity",

            order_items_filename,

            "order_id",

            "Order items reference orders that do not exist",

            orphan_items,

            percentage,

            "High",

            "Investigate orphan order-item records before warehouse loading."

        )


    if orders is not None:

        orders_without_items = int(

            (
                ~orders[
                    "order_id"
                ].isin(item_order_ids)
            ).sum()

        )

        if orders_without_items > 0:

            percentage = (

                orders_without_items /
                len(orders) *
                100

            )

            add_finding(

                findings,

                "Referential Integrity",

                orders_filename,

                "order_id",

                "Orders without corresponding order-item records",

                orders_without_items,

                percentage,

                "Review",

                "Review against order status. Cancelled or unavailable orders may legitimately have no order-item records."

            )


if payments is not None:

    payment_order_ids = set(

        payments[
            "order_id"
        ].dropna()

    )

    orphan_payments = int(

        (
            ~payments[
                "order_id"
            ].isin(order_ids)
        ).sum()

    )

    if orphan_payments > 0:

        percentage = (

            orphan_payments /
            len(payments) *
            100

        )

        add_finding(

            findings,

            "Referential Integrity",

            payments_filename,

            "order_id",

            "Payments reference orders that do not exist",

            orphan_payments,

            percentage,

            "High",

            "Investigate orphan payment records before warehouse loading."

        )


    if orders is not None:

        orders_without_payments = int(

            (
                ~orders[
                    "order_id"
                ].isin(payment_order_ids)
            ).sum()

        )

        if orders_without_payments > 0:

            percentage = (

                orders_without_payments /
                len(orders) *
                100

            )

            add_finding(

                findings,

                "Referential Integrity",

                orders_filename,

                "order_id",

                "Orders without corresponding payment records",

                orders_without_payments,

                percentage,

                "Review",

                "Investigate missing payment records and determine whether order status explains the absence."

            )


if order_items is not None and products is not None:

    product_ids = set(

        products[
            "product_id"
        ].dropna()

    )

    orphan_products = int(

        (
            ~order_items[
                "product_id"
            ].isin(product_ids)
        ).sum()

    )

    if orphan_products > 0:

        percentage = (

            orphan_products /
            len(order_items) *
            100

        )

        add_finding(

            findings,

            "Referential Integrity",

            order_items_filename,

            "product_id",

            "Order items reference products that do not exist",

            orphan_products,

            percentage,

            "High",

            "Investigate orphan product references before warehouse loading."

        )


if order_items is not None and sellers is not None:

    seller_ids = set(

        sellers[
            "seller_id"
        ].dropna()

    )

    orphan_sellers = int(

        (
            ~order_items[
                "seller_id"
            ].isin(seller_ids)
        ).sum()

    )

    if orphan_sellers > 0:

        percentage = (

            orphan_sellers /
            len(order_items) *
            100

        )

        add_finding(

            findings,

            "Referential Integrity",

            order_items_filename,

            "seller_id",

            "Order items reference sellers that do not exist",

            orphan_sellers,

            percentage,

            "High",

            "Investigate orphan seller references before warehouse loading."

        )


print(
    "Referential integrity checks completed."
)


print("\n" + "-" * 75)
print("6. ORDER DATE CONSISTENCY CHECKS")
print("-" * 75)


if orders is not None:

    date_columns = [

        "order_purchase_timestamp",

        "order_approved_at",

        "order_delivered_carrier_date",

        "order_delivered_customer_date",

        "order_estimated_delivery_date"

    ]

    for column in date_columns:

        if column in orders.columns:

            orders[column] = pd.to_datetime(

                orders[column],

                errors="coerce"

            )


    if (

        "order_purchase_timestamp" in orders.columns
        and
        "order_approved_at" in orders.columns

    ):

        invalid_sequence = (

            orders[
                "order_approved_at"
            ].notna()

            &

            orders[
                "order_purchase_timestamp"
            ].notna()

            &

            (
                orders[
                    "order_approved_at"
                ]

                <

                orders[
                    "order_purchase_timestamp"
                ]
            )

        )

        count = int(
            invalid_sequence.sum()
        )

        if count > 0:

            percentage = (
                count /
                len(orders) *
                100
            )

            add_finding(

                findings,

                "Consistency",

                orders_filename,

                "order_approved_at",

                "Approval timestamp occurs before purchase timestamp",

                count,

                percentage,

                "High",

                "Investigate timestamp ordering and correct invalid date relationships."

            )


    if (

        "order_delivered_carrier_date"
        in orders.columns

        and

        "order_purchase_timestamp"
        in orders.columns

    ):

        invalid_sequence = (

            orders[
                "order_delivered_carrier_date"
            ].notna()

            &

            orders[
                "order_purchase_timestamp"
            ].notna()

            &

            (
                orders[
                    "order_delivered_carrier_date"
                ]

                <

                orders[
                    "order_purchase_timestamp"
                ]
            )

        )

        count = int(
            invalid_sequence.sum()
        )

        if count > 0:

            percentage = (
                count /
                len(orders) *
                100
            )

            add_finding(

                findings,

                "Consistency",

                orders_filename,

                "order_delivered_carrier_date",

                "Carrier delivery occurs before purchase",

                count,

                percentage,

                "High",

                "Investigate invalid order timestamp sequences."

            )


    if (

        "order_delivered_customer_date"
        in orders.columns

        and

        "order_delivered_carrier_date"
        in orders.columns

    ):

        invalid_sequence = (

            orders[
                "order_delivered_customer_date"
            ].notna()

            &

            orders[
                "order_delivered_carrier_date"
            ].notna()

            &

            (
                orders[
                    "order_delivered_customer_date"
                ]

                <

                orders[
                    "order_delivered_carrier_date"
                ]
            )

        )

        count = int(
            invalid_sequence.sum()
        )

        if count > 0:

            percentage = (
                count /
                len(orders) *
                100
            )

            add_finding(

                findings,

                "Consistency",

                orders_filename,

                "order_delivered_customer_date",

                "Customer delivery occurs before carrier delivery",

                count,

                percentage,

                "High",

                "Investigate invalid delivery timestamp sequences."

            )


    if (

        "order_estimated_delivery_date"
        in orders.columns

        and

        "order_purchase_timestamp"
        in orders.columns

    ):

        purchase_date = orders[
            "order_purchase_timestamp"
        ].dt.normalize()

        estimated_date = orders[
            "order_estimated_delivery_date"
        ].dt.normalize()

        invalid_sequence = (

            estimated_date.notna()

            &

            purchase_date.notna()

            &

            (
                estimated_date
                <
                purchase_date
            )

        )

        count = int(
            invalid_sequence.sum()
        )

        if count > 0:

            percentage = (
                count /
                len(orders) *
                100
            )

            add_finding(

                findings,

                "Consistency",

                orders_filename,

                "order_estimated_delivery_date",

                "Estimated delivery occurs before purchase",

                count,

                percentage,

                "Medium",

                "Investigate invalid estimated delivery dates."

            )


print(
    "Order date consistency checks completed."
)


print("\n" + "-" * 75)
print("7. GEOLOCATION CHECKS")
print("-" * 75)


geolocation_filename = (
    "olist_geolocation_dataset.csv"
)


if geolocation_filename in datasets:

    df = datasets[
        geolocation_filename
    ]

    duplicate_geo = int(
        df.duplicated().sum()
    )

    if duplicate_geo > 0:

        percentage = (

            duplicate_geo /
            len(df) *
            100

        )

        add_finding(

            findings,

            "Uniqueness Observation",

            geolocation_filename,

            "Full Record",

            "Repeated geolocation records detected",

            duplicate_geo,

            percentage,

            "Review",

            "Do not automatically delete. Repeated geographic coordinates can legitimately occur for multiple postal-code or geographic records."

        )

    print(
        "Negative latitude and longitude values are not classified as errors."
    )

    print(
        "Reason: negative geographic coordinates are valid."
    )


print("\n" + "-" * 75)
print("8. FINAL FINDINGS PREPARATION")
print("-" * 75)


findings_df = pd.DataFrame(
    findings
)


if findings_df.empty:

    findings_df = pd.DataFrame(

        columns=[

            "Data Quality Dimension",

            "Dataset",

            "Column",

            "Issue",

            "Count",

            "Percentage",

            "Severity",

            "Recommended Treatment"

        ]

    )


severity_order = {

    "High": 1,

    "Medium": 2,

    "Low": 3,

    "Review": 4

}


if not findings_df.empty:

    findings_df[
        "Severity Order"
    ] = (

        findings_df[
            "Severity"
        ]

        .map(severity_order)

        .fillna(5)

    )

    findings_df = findings_df.sort_values(

        [

            "Severity Order",

            "Data Quality Dimension",

            "Dataset"

        ]

    )

    findings_df = findings_df.drop(

        columns=[
            "Severity Order"
        ]

    )


summary_df = pd.DataFrame(
    summary_rows
)


print(
    f"Total quality findings: {len(findings_df)}"
)


print("\n" + "-" * 75)
print("9. GENERATING EXCEL QUALITY FINDINGS")
print("-" * 75)


with pd.ExcelWriter(

    EXCEL_OUTPUT,

    engine="openpyxl"

) as writer:

    summary_df.to_excel(

        writer,

        sheet_name="Dataset_Summary",

        index=False

    )


    findings_df.to_excel(

        writer,

        sheet_name="Quality_Findings",

        index=False

    )


    completeness_df = findings_df[
        findings_df[
            "Data Quality Dimension"
        ].isin(

            [

                "Completeness",

                "Product Completeness"

            ]

        )

    ]

    completeness_df.to_excel(

        writer,

        sheet_name="Completeness",

        index=False

    )


    uniqueness_df = findings_df[
        findings_df[
            "Data Quality Dimension"
        ].isin(

            [

                "Uniqueness",

                "Uniqueness Observation"

            ]

        )

    ]

    uniqueness_df.to_excel(

        writer,

        sheet_name="Uniqueness",

        index=False

    )


    validity_df = findings_df[
        findings_df[
            "Data Quality Dimension"
        ] == "Validity"

    ]

    validity_df.to_excel(

        writer,

        sheet_name="Validity",

        index=False

    )


    referential_df = findings_df[
        findings_df[
            "Data Quality Dimension"
        ] == "Referential Integrity"

    ]

    referential_df.to_excel(

        writer,

        sheet_name="Referential_Integrity",

        index=False

    )


    consistency_df = findings_df[
        findings_df[
            "Data Quality Dimension"
        ] == "Consistency"

    ]

    consistency_df.to_excel(

        writer,

        sheet_name="Consistency",

        index=False

    )


print(
    "Excel findings report created."
)


print("\n" + "-" * 75)
print("10. GENERATING WORD DATA QUALITY REPORT")
print("-" * 75)


document = Document()


section = document.sections[0]

section.top_margin = Inches(0.7)

section.bottom_margin = Inches(0.7)

section.left_margin = Inches(0.7)

section.right_margin = Inches(0.7)


title = document.add_heading(

    "Sprint 2 – Data Quality Report",

    level=0

)

title.alignment = (
    WD_ALIGN_PARAGRAPH.CENTER
)


subtitle = document.add_paragraph(

    "Enterprise Sales Analytics Data Platform (ESADP)"

)

subtitle.alignment = (
    WD_ALIGN_PARAGRAPH.CENTER
)


document.add_paragraph(

    "This report presents the data quality assessment of the Olist "
    "source datasets used for the Enterprise Sales Analytics Data "
    "Platform. The assessment evaluates completeness, uniqueness, "
    "validity, referential integrity, consistency, and product data quality."

)


document.add_heading(

    "1. Data Quality Assessment Scope",

    level=1

)


document.add_paragraph(

    "The assessment is performed on the original Olist source CSV "
    "datasets before warehouse transformation. The objective is to "
    "identify data quality issues that should be considered during "
    "data cleaning, standardization, and warehouse loading."

)


document.add_heading(

    "2. Dataset Summary",

    level=1

)


total_rows = int(
    summary_df[
        "Rows"
    ].sum()
)

total_columns = int(
    summary_df[
        "Columns"
    ].sum()
)

total_missing = int(
    summary_df[
        "Missing Cells"
    ].sum()
)

total_duplicates = int(
    summary_df[
        "Duplicate Full Rows"
    ].sum()
)


summary_report = pd.DataFrame({

    "Metric": [

        "Datasets analyzed",

        "Total rows",

        "Total columns",

        "Total missing cells",

        "Total duplicate full rows",

        "Total quality findings"

    ],

    "Value": [

        len(summary_df),

        total_rows,

        total_columns,

        total_missing,

        total_duplicates,

        len(findings_df)

    ]

})


add_table(

    document,

    summary_report,

    max_rows=20

)


document.add_heading(

    "3. Completeness Assessment",

    level=1

)


completeness_report = findings_df[

    findings_df[
        "Data Quality Dimension"
    ].isin(

        [

            "Completeness",

            "Product Completeness"

        ]

    )

]


if completeness_report.empty:

    document.add_paragraph(

        "No completeness issues were identified."

    )

else:

    document.add_paragraph(

        f"{len(completeness_report)} completeness findings were identified. "
        "Severity is based on missing-value percentage, while final "
        "treatment depends on the business meaning of each field."

    )

    add_table(

        document,

        completeness_report[

            [

                "Dataset",

                "Column",

                "Issue",

                "Count",

                "Percentage",

                "Severity",

                "Recommended Treatment"

            ]

        ],

        max_rows=40

    )


document.add_heading(

    "4. Uniqueness Assessment",

    level=1

)


uniqueness_report = findings_df[

    findings_df[
        "Data Quality Dimension"
    ].isin(

        [

            "Uniqueness",

            "Uniqueness Observation"

        ]

    )

]


if uniqueness_report.empty:

    document.add_paragraph(

        "No key uniqueness issues were identified."

    )

else:

    add_table(

        document,

        uniqueness_report[

            [

                "Dataset",

                "Column",

                "Issue",

                "Count",

                "Percentage",

                "Severity",

                "Recommended Treatment"

            ]

        ],

        max_rows=40

    )


document.add_paragraph(

    "For the Olist order-item dataset, uniqueness is evaluated using "
    "the composite key order_id + order_item_id because order_item_id "
    "alone can legitimately repeat across different orders."

)


document.add_paragraph(

    "Repeated customer_unique_id values are not automatically treated "
    "as duplicate customers because the same customer can be associated "
    "with multiple transactions."

)


document.add_heading(

    "5. Numeric and Validity Assessment",

    level=1

)


validity_report = findings_df[

    findings_df[
        "Data Quality Dimension"
    ] == "Validity"

]


if validity_report.empty:

    document.add_paragraph(

        "No invalid numeric or product measurement values were identified."

    )

else:

    add_table(

        document,

        validity_report[

            [

                "Dataset",

                "Column",

                "Issue",

                "Count",

                "Percentage",

                "Severity",

                "Recommended Treatment"

            ]

        ],

        max_rows=40

    )


document.add_paragraph(

    "Geolocation latitude and longitude values were not classified as "
    "invalid simply because they are negative. Negative geographic "
    "coordinates are valid."

)


document.add_heading(

    "6. Referential Integrity Assessment",

    level=1

)


referential_report = findings_df[

    findings_df[
        "Data Quality Dimension"
    ] == "Referential Integrity"

]


if referential_report.empty:

    document.add_paragraph(

        "No referential integrity issues were identified."

    )

else:

    document.add_paragraph(

        "Referential integrity checks validate relationships between "
        "orders, order items, payments, products, and sellers."

    )

    add_table(

        document,

        referential_report[

            [

                "Dataset",

                "Column",

                "Issue",

                "Count",

                "Percentage",

                "Severity",

                "Recommended Treatment"

            ]

        ],

        max_rows=40

    )


document.add_heading(

    "7. Order Date Consistency Assessment",

    level=1

)


consistency_report = findings_df[

    findings_df[
        "Data Quality Dimension"
    ] == "Consistency"

]


if consistency_report.empty:

    document.add_paragraph(

        "No invalid order timestamp sequences were identified."

    )

else:

    add_table(

        document,

        consistency_report[

            [

                "Dataset",

                "Column",

                "Issue",

                "Count",

                "Percentage",

                "Severity",

                "Recommended Treatment"

            ]

        ],

        max_rows=40

    )


document.add_heading(

    "8. Product Data Quality Assessment",

    level=1

)


product_report = findings_df[

    findings_df[
        "Data Quality Dimension"
    ] == "Product Completeness"

]


if product_report.empty:

    document.add_paragraph(

        "No product completeness issues were identified."

    )

else:

    add_table(

        document,

        product_report[

            [

                "Dataset",

                "Column",

                "Issue",

                "Count",

                "Percentage",

                "Severity",

                "Recommended Treatment"

            ]

        ],

        max_rows=40

    )


document.add_heading(

    "9. Data Cleaning Recommendations",

    level=1

)


recommendations = [

    "Investigate high-percentage missing attributes and determine whether NULL retention, imputation, or default values are appropriate.",

    "Validate duplicate primary and business identifiers before warehouse loading.",

    "Use the composite order_id and order_item_id key when validating order-item uniqueness.",

    "Do not remove repeated geolocation records without understanding their geographic meaning.",

    "Validate negative prices, freight values, and product measurements.",

    "Validate referential relationships between orders, order items, payments, products, and sellers.",

    "Investigate orders without item-level or payment records using order status and business context.",

    "Validate chronological consistency among purchase, approval, carrier delivery, customer delivery, and estimated delivery timestamps.",

    "Standardize product master attributes and document treatment of missing product information.",

    "Document all cleaning and transformation rules so that the warehouse loading process is reproducible."

]


for recommendation in recommendations:

    document.add_paragraph(

        recommendation,

        style="List Bullet"

    )


document.add_heading(

    "10. Conclusion",

    level=1

)


document.add_paragraph(

    "The Olist source datasets were evaluated across completeness, "
    "uniqueness, validity, referential integrity, consistency, and "
    "product data quality dimensions. The resulting findings provide "
    "the baseline for Sprint 2 data cleaning and standardization. "
    "The identified issues should be investigated and addressed "
    "through documented transformation rules before the data is "
    "loaded into the analytical warehouse."

)


document.add_paragraph()


document.add_paragraph(

    "Source: Olist Kaggle source CSV datasets."

)


document.save(

    WORD_OUTPUT

)


print("\n" + "=" * 75)
print("DATA QUALITY REPORT COMPLETED")
print("=" * 75)


print(
    f"\nDatasets analyzed : {len(summary_df)}"
)

print(
    f"Total rows       : {total_rows:,}"
)

print(
    f"Total columns    : {total_columns:,}"
)

print(
    f"Missing cells    : {total_missing:,}"
)

print(
    f"Duplicate rows   : {total_duplicates:,}"
)

print(
    f"Quality findings : {len(findings_df)}"
)

print(
    "\nWord report:"
)

print(
    os.path.abspath(WORD_OUTPUT)
)

print(
    "\nExcel findings:"
)

print(
    os.path.abspath(EXCEL_OUTPUT)
)